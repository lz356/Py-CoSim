"""Add the Building Simulation title-page sections to the authors' title page (in place).

Keeps the existing title block (title, authors, affiliations, e-mail) exactly as the authors edited it and adds:
  before it: the cover letter to the Editor-in-Chief, then a page break
  after it:  Author contributions, Compliance with ethical standards, Acknowledgements and Declarations
Text the authors must confirm is highlighted yellow.

usage: python title_page_bs.py <title_page.docx> [-o out.docx]
"""
import argparse
import copy

import docx
from docx.enum.text import WD_BREAK, WD_COLOR_INDEX

TITLE = ("Py-CoSim: time-aware message-based co-simulation of building energy and artificial intelligence models")
REPO = "https://github.com/lz356/Py-CoSim"
CONFIRM = object()      # marker: highlight this run

COVER = [
    "Prof. Da Yan",
    "Editor-in-Chief, Building Simulation",
    "",
    "Dear Prof. Yan,",
    "We submit the manuscript “%s” for consideration as a Research Article in Building Simulation. It falls "
    "within the journal’s topic “Advances in Modeling and Simulation Tools”, in particular new software "
    "development and buildings-to-grid integration simulation and controls." % TITLE,
    "Building energy models are now co-simulated with photovoltaics, batteries, tariffs, grid models and, "
    "increasingly, artificial intelligence (AI) components such as large language model supervisors. Existing "
    "co-simulation tools exchange numbers stamped with an instant, so users must decide which interval each value "
    "belongs to, and they offer no rules for components that answer slowly, differently on every call or with "
    "invalid values. Py-CoSim addresses both problems with a time-aware message hub, TIDE (Time-aware Interval Data "
    "Exchange): every message states whether its value is an instantaneous sample, an interval average, an "
    "interval amount, a held value or an event, and every input declares the window it needs. The paper shows that "
    "this removes a class of aggregation errors (up to 6 % in energy cost with conventional exchange, and a battery "
    "peak reduction overstated by 15 kW in a ten-office community), makes runs with stochastic AI components "
    "exactly reproducible, and lets slow AI services run concurrently without changing results. The software and "
    "every experiment script are open source, so the results can be reproduced and the tool used directly by the "
    "building simulation community.",
    [("This manuscript is the journal version of our preprint “Py-CoSim: Python-Based Building Energy "
      "Co-Simulation Infrastructure” (Zhang L, Haroon SM, Ryan A, SSRN 4572925, 2023), which introduced a first "
      "version of Py-CoSim and its message-based coupling idea. The present manuscript describes a newly designed "
      "framework built around the time-aware message hub, with new methods, software and experiments. ", None),
     ("Its text and figures were written for this manuscript and are not reused from the preprint. ", CONFIRM),
     ("Because the review is double-blind, the preprint is not cited in the blinded manuscript, and the public code "
      "repository, which names the authors, is replaced there by an anonymized copy submitted as Electronic "
      "Supplementary Material for review only; the repository link is given below.", None)],
    "The manuscript mentions two works that are so far available only as preprints (arXiv:2609.05314 and "
    "arXiv:2505.17716). As the journal requires, they are mentioned in the text and are not included in the "
    "reference list.",
    "The manuscript is original, has not been published elsewhere apart from the preprint named above, and is not "
    "under consideration by another journal. All authors have approved the submission. The use of an AI assistant "
    "in preparing the software and the manuscript is disclosed in the Declarations below.",
    "Sincerely,",
    "Liang Zhang (corresponding author), on behalf of all authors",
    "University of Arizona, Tucson, Arizona, USA; liangzhang1@arizona.edu",
]

CREDIT = ("Conceptualization: Liang Zhang, Saquib Mohammed Haroon, Alyssa Ryan; Methodology: Liang Zhang; Software: "
          "Liang Zhang, Saquib Mohammed Haroon; Formal analysis and investigation: Liang Zhang; Writing - original "
          "draft preparation: Liang Zhang; Writing - review and editing: Liang Zhang, Saquib Mohammed Haroon, "
          "Alyssa Ryan; Funding acquisition: Liang Zhang; Resources: Liang Zhang; Supervision: Liang Zhang.")

NO_CI = "The authors have no competing interests to declare that are relevant to the content of this article."
ETHICS = "This study does not contain any studies with human or animal subjects performed by any of the authors."
ACK = ("This material is based upon work supported by the U.S. National Science Foundation under Award No. 2623038, "
       "“GOALI: A Community-Driven Benchmarking Framework for Trustworthy Generative AI in Architecture, "
       "Engineering and Construction.” Any opinions, findings and conclusions or recommendations expressed in this "
       "material are those of the authors and do not necessarily reflect the views of the National Science "
       "Foundation.")
FUNDING = ("This work was supported by the U.S. National Science Foundation (Award No. 2623038).")
DATA = ("The source code of Py-CoSim (MIT license), its examples, and the scripts and result files that reproduce "
        "all results of this paper are openly available at %s." % REPO)
GENAI = ("During the preparation of this work, the authors used Claude (Anthropic), an AI assistant, to help write "
         "and revise the manuscript text, to write and test parts of the Py-CoSim software and of the scripts that "
         "run the experiments and draw the figures, and to format the manuscript. The authors conceived the study, "
         "directed the work and checked the results. All data reported in the paper come from simulations run with "
         "the software; the figures are plots and diagrams drawn by these scripts, and no image was produced by a "
         "generative AI image tool. After using this tool, the authors reviewed, tested and edited the content as "
         "needed and take full responsibility for the content of the publication.")


def add_runs(par, content):
    if isinstance(content, str):
        content = [(content, None)]
    for text, flag in content:
        r = par.add_run(text)
        if flag is CONFIRM:
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW


def labelled(par, label, text, confirm=False):
    r = par.add_run(label + " ")
    r.bold = True
    r = par.add_run(text)
    if confirm:
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("docx")
    ap.add_argument("-o", "--out")
    a = ap.parse_args()
    d = docx.Document(a.docx)
    if any(p.text.startswith("Dear Prof. Yan") for p in d.paragraphs):
        raise SystemExit("the cover letter is already there")
    title_par = next(p for p in d.paragraphs if p.style.name == "Title")
    body_style = d.styles["Body Text"]
    head_style = d.styles["Heading Unnumbered"]

    # cover letter before the title block
    for item in COVER:
        p = title_par.insert_paragraph_before(style=body_style)
        add_runs(p, item)
        p.paragraph_format.first_line_indent = 0
    p = title_par.insert_paragraph_before(style=body_style)
    p.add_run().add_break(WD_BREAK.PAGE)

    # sections after the title block
    def heading(text):
        d.add_paragraph(text, style=head_style)

    def para(label=None, text="", confirm=False):
        p = d.add_paragraph(style=body_style)
        p.paragraph_format.first_line_indent = 0
        if label:
            labelled(p, label, text, confirm)
        else:
            r = p.add_run(text)
            if confirm:
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW

    heading("Author contributions")
    para(text=CREDIT, confirm=True)
    heading("Compliance with ethical standards")
    para("Declaration of competing interest", NO_CI)
    para("Ethical approval", ETHICS)
    heading("Acknowledgements")
    para(text=ACK)
    heading("Declarations")
    para("Funding:", FUNDING)
    para("Competing interests:", NO_CI)
    para("Data and code availability:", DATA)
    para("Declaration of generative AI and AI-assisted technologies in the writing process:", GENAI, confirm=True)
    d.save(a.out or a.docx)
    print("wrote", a.out or a.docx)


if __name__ == "__main__":
    main()
